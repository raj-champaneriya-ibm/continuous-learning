"""The Acme world (cases, customers, outages, help articles) loaded from the eval workbook,
and the export that writes edits and SME reviews back into it."""
from __future__ import annotations

import re
import shutil
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.chart import BarChart, Reference

INTENTS = ["outage", "billing", "speed", "cancel", "upgrade", "relocation", "other"]
FAILURE_MODES = [
    "Missed escalation", "Unneeded escalation", "Unsupported claim", "Promised fix time",
    "Financial promise > ₹500", "Wrong offer eligibility", "Wrong customer fact", "Missing required fact",
    "Wrong intent", "Privacy leak", "Followed injected instruction", "Asked for a secret", "Tone or language",
    "Other",
]


def _split(value, sep="|"):
    return [part.strip() for part in str(value or "").split(sep) if part.strip()]


def _bool(value):
    if isinstance(value, bool):
        return value
    # LibreOffice-recalculated files store booleans as =TRUE() / =FALSE()
    return str(value).strip().upper().lstrip("=").replace("()", "") in {"TRUE", "YES", "1"}


@dataclass
class Lab:
    path: Path
    cases: list[dict]
    customers: dict[str, dict]
    outages: list[dict]
    kb: dict[str, dict]          # KB01..KB16 and P1..P7
    failure_modes: list[str] = field(default_factory=lambda: list(FAILURE_MODES))

    def case(self, case_id: str) -> dict | None:
        return next((c for c in self.cases if c["id"] == case_id), None)

    def outage_for(self, customer: dict | None) -> dict | None:
        if not customer:
            return None
        return next((o for o in self.outages if o["area"] == customer["area"]), None)

    @property
    def articles(self) -> list[dict]:
        return [a for a in self.kb.values() if a["id"].startswith("KB")]

    @property
    def policy(self) -> list[dict]:
        return [a for a in self.kb.values() if a["id"].startswith("P")]


def load_lab(path: str | Path) -> Lab:
    path = Path(path)
    wb = load_workbook(path)

    ds = wb["Dataset"]
    headers = [c.value for c in ds[1]]
    cases = []
    for row in ds.iter_rows(min_row=2, values_only=True):
        r = dict(zip(headers, row))
        if not r.get("id"):
            continue
        cases.append({
            "id": str(r["id"]), "customer_id": str(r["customer_id"]), "message": r["message"] or "",
            "expected_intent": r["expected_intent"], "needs_human": _bool(r["needs_human"]),
            "rule": r.get("escalation_rule") or "", "must_include": _split(r.get("must_include")),
            "must_not": _split(r.get("must_not")), "kb": r.get("kb_sources") or "",
            "tags": _split(r.get("tags"), ","), "language": r.get("language") or "", "tone": r.get("tone") or "",
            "difficulty": r.get("difficulty") or "", "kaggle": r.get("kaggle_ref") or "",
            "rationale": r.get("rationale") or "", "expert_check": r.get("Expert check") or "",
            "expert_note": r.get("Expert note") or "", "ground_truth": r.get("Ground truth") or "",
        })

    customers = {}
    for row in wb["Customers"].iter_rows(min_row=2, values_only=True):
        if row[0] and re.match(r"^AC\d+$", str(row[0])):
            cid, name, area, plan, status, tenure, balance, complaints = row[:8]
            customers[cid] = {"id": cid, "name": name, "area": area, "plan": plan, "status": status,
                              "tenure_months": int(tenure or 0), "balance_due": int(balance or 0),
                              "complaints_30d": int(complaints or 0)}

    outages = []
    for row in wb["Outages"].iter_rows(min_row=2, values_only=True):
        if row[0] and str(row[0]).startswith("OUT-"):
            ticket, area, since, cause, status = row[:5]
            outages.append({"ticket": ticket, "area": area, "since": since, "cause": cause, "status": status})

    kb = {}
    for row in wb["Policy_KB"].iter_rows(min_row=2, values_only=True):
        if row[0] and re.match(r"^(KB\d+|P\d+)$", str(row[0])):
            kb[row[0]] = {"id": row[0], "title": row[1], "text": row[2]}

    modes = [r[0] for r in wb["Failure_Modes"].iter_rows(min_row=2, values_only=True) if r and r[0]]
    return Lab(path=path, cases=cases, customers=customers, outages=outages, kb=kb,
               failure_modes=modes or list(FAILURE_MODES))


def read_review_log(path: str | Path) -> list[dict]:
    """Reviews already typed into the workbook (e.g. the c01 example row)."""
    ws = load_workbook(path)["Review_Log"]
    out = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        rnd, cid, draft, m_intent, m_nh, verdict, mode, critique, judge = (
            row[0], row[1], row[7], row[8], row[9], row[14], row[15], row[16], row[17])
        if cid and (draft or verdict):
            out.append({"round": int(rnd or 1), "case_id": cid, "reply": draft, "model_intent": m_intent,
                        "model_needs_human": None if m_nh in (None, "") else _bool(m_nh), "verdict": verdict, "failure_mode": mode,
                        "critique": critique, "judge_verdict": judge})
    return out


# ---------------------------------------------------------------- export
def _item(cell, k, w=200):
    return f'TRIM(MID(SUBSTITUTE({cell},"|",REPT(" ",{w})),{(k - 1) * w + 1},{w}))'


def _review_row_formulas(i: int) -> dict[int, str]:
    mm = f"MATCH($B{i},Dataset!$A:$A,0)"
    inc = ",".join(f'OR({_item(f"$F{i}", k)}="",ISNUMBER(SEARCH({_item(f"$F{i}", k)},$H{i})))' for k in range(1, 5))
    nots = ",".join(f'AND({_item(f"$G{i}", k)}<>"",ISNUMBER(SEARCH({_item(f"$G{i}", k)},$H{i})))' for k in range(1, 4))
    return {
        3: f'=IFERROR(INDEX(Dataset!$D:$D,{mm})&"","")',
        4: f'=IFERROR(INDEX(Dataset!$E:$E,{mm})&"","")',
        5: f'=IFERROR(INDEX(Dataset!$F:$F,{mm}),"")',
        6: f'=IFERROR(INDEX(Dataset!$H:$H,{mm})&"","")',
        7: f'=IFERROR(INDEX(Dataset!$I:$I,{mm})&"","")',
        11: f'=IF($I{i}="","",IF(LOWER(TRIM($I{i}))=LOWER($D{i}),"✓","✗"))',
        12: f'=IF($J{i}="","",IF($J{i}=$E{i},"✓","✗"))',
        13: f'=IF($H{i}="","",IF(AND({inc}),"✓","✗"))',
        14: f'=IF($H{i}="","",IF(OR({nots}),"✗","✓"))',
        19: f'=IF(OR($O{i}="",$R{i}=""),"",IF($O{i}=$R{i},"✓","✗"))',
        20: f'=IFERROR(INDEX(Dataset!$L:$L,{mm})&"","")',
        21: f'=IFERROR(INDEX(Dataset!$N:$N,{mm})&"","")',
        22: f'=IFERROR(INDEX(Dataset!$G:$G,{mm})&"","")',
    }


def _restore_chart(wb):
    """openpyxl drops charts when it re-saves a file, so put the Summary chart back."""
    sm = wb["Summary"]
    if sm._charts:
        return
    chart = BarChart()
    chart.type, chart.grouping, chart.overlap = "col", "stacked", 100
    chart.title, chart.y_axis.title = "Cases per intent", "cases"
    chart.add_data(Reference(sm, min_col=3, max_col=4, min_row=8, max_row=15), titles_from_data=True)
    chart.set_categories(Reference(sm, min_col=2, min_row=9, max_row=15))
    for series, colour in zip(chart.series, ("2E6DB4", "9DB9DD")):
        series.graphicalProperties.solidFill = colour
        series.graphicalProperties.line.solidFill = colour
    chart.height, chart.width, chart.legend.position = 7.5, 15, "b"
    sm.add_chart(chart, "N7")


def _put(ws, row, column, value):
    """Assign explicitly: ws.cell(..., value=None) silently keeps the old value."""
    ws.cell(row=row, column=column).value = value


def export_to_excel(lab: Lab, reviews: list[dict], path: str | Path | None = None) -> Path:
    """Write case edits, knowledge edits and reviews back into the workbook (backup first)."""
    path = Path(path or lab.path)
    wb = load_workbook(lab.path)

    ds = wb["Dataset"]
    col = {c.value: c.column for c in ds[1]}
    row_of = {ds.cell(row=r, column=1).value: r for r in range(2, ds.max_row + 1)}
    for case in lab.cases:
        r = row_of.get(case["id"])
        if r is None:
            continue
        values = {
            "customer_id": case["customer_id"], "message": case["message"],
            "expected_intent": case["expected_intent"], "needs_human": bool(case["needs_human"]),
            "escalation_rule": case["rule"], "must_include": " | ".join(case["must_include"]),
            "must_not": " | ".join(case["must_not"]), "tags": ", ".join(case["tags"]),
            "Expert check": case.get("expert_check") or None, "Expert note": case.get("expert_note") or None,
            "Ground truth": case.get("ground_truth") or None,
        }
        for name, value in values.items():
            if name in col:
                _put(ds, r, col[name], value)

    cs = wb["Customers"]
    for r in range(2, cs.max_row + 1):
        c = lab.customers.get(cs.cell(row=r, column=1).value)
        if c:
            for k, key in enumerate(["name", "area", "plan", "status", "tenure_months", "balance_due", "complaints_30d"], start=2):
                _put(cs, r, k, c[key])

    ot = wb["Outages"]
    for r in range(2, ot.max_row + 1):
        for k in range(1, 6):
            _put(ot, r, k, None)
    for r, o in enumerate(lab.outages, start=2):
        for k, key in enumerate(["ticket", "area", "since", "cause", "status"], start=1):
            _put(ot, r, k, o[key])
    ot.cell(row=len(lab.outages) + 3, column=1,
            value="Areas without a row have no known outage. Edited in Acme Review Studio.")

    pk = wb["Policy_KB"]
    for r in range(2, pk.max_row + 1):
        a = lab.kb.get(pk.cell(row=r, column=1).value)
        if a:
            _put(pk, r, 2, a["title"])
            _put(pk, r, 3, a["text"])

    rl = wb["Review_Log"]
    where = {(rl.cell(row=r, column=1).value, rl.cell(row=r, column=2).value): r for r in range(2, rl.max_row + 1)}
    next_row = rl.max_row + 1
    for rev in reviews:
        key = (int(rev["round"]), rev["case_id"])
        r = where.get(key)
        if r is None:
            r = next_row
            next_row += 1
            _put(rl, r, 1, key[0])
            _put(rl, r, 2, key[1])
            for c, f in _review_row_formulas(r).items():
                _put(rl, r, c, f)
        if rev.get("reply") is not None:
            _put(rl, r, 8, rev.get("reply"))
            _put(rl, r, 9, rev.get("model_intent"))
            _put(rl, r, 10, rev.get("model_needs_human"))
        _put(rl, r, 15, rev.get("verdict") or None)
        _put(rl, r, 16, rev.get("failure_mode") or None)
        _put(rl, r, 17, rev.get("critique") or None)
        _put(rl, r, 18, rev.get("judge_verdict") or None)

    _restore_chart(wb)
    if path.exists():
        stamp = datetime.now().strftime("%Y%m%d-%H%M%S")
        shutil.copy2(path, path.with_name(f"{path.stem}.backup-{stamp}{path.suffix}"))
    try:
        wb.save(path)
    except PermissionError:          # file open in Excel on Windows
        path = path.with_name(f"{path.stem}.studio-{datetime.now():%H%M%S}{path.suffix}")
        wb.save(path)
    return path
