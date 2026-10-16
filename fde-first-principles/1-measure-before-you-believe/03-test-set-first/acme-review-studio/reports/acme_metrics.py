"""One definition of every number the reports show.

Every page imports this module, so "pass rate" or "missed escalation" means the same thing in the
evaluation report, the dataset card and the review dashboard. It only reads saved files; it never
calls a model or changes a review.

Inputs (override with environment variables):
    ACME_XLSX   the eval workbook             default: ../data/Acme_Fibre_Grounded_Draft_Eval.xlsx
    ACME_RUNS   Review Studio's saved runs     default: ../runs/reviews.json
"""
from __future__ import annotations

import hashlib
import json
import math
import os
import sys
from dataclasses import dataclass, field
from pathlib import Path

import pandas as pd
import yaml
from openpyxl import load_workbook

HERE = Path(__file__).resolve().parent
ROOT = HERE.parent                                   # acme-review-studio/
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from acme.checks import eval_checks  # noqa: E402
from acme.lab import INTENTS, load_lab  # noqa: E402

WORKBOOK = Path(os.environ.get("ACME_XLSX", ROOT / "data" / "Acme_Fibre_Grounded_Draft_Eval.xlsx"))
RUNS = Path(os.environ.get("ACME_RUNS", ROOT / "runs" / "reviews.json"))

# Colours: validated categorical pair, one sequential hue, status colours only next to an icon + label.
BLUE, ORANGE, GRAY, INK, MUTED = "#2a78d6", "#eb6834", "#b4b2a9", "#1A2433", "#5B6B80"
GOOD, CRITICAL = "#0ca30c", "#d03b3b"

LAYERS = ["Model flag", "Code guard", "Step-6 check", "Stopped early"]


# ------------------------------------------------------------------ small helpers
def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest() if path.exists() else "missing"


def wilson(k: int, n: int, z: float = 1.96) -> tuple[float, float]:
    """95% Wilson interval for k successes out of n. Behaves well for small n and rates near 0 or 1."""
    if n == 0:
        return (math.nan, math.nan)
    p = k / n
    denom = 1 + z * z / n
    centre = (p + z * z / (2 * n)) / denom
    half = z * math.sqrt(p * (1 - p) / n + z * z / (4 * n * n)) / denom
    return (max(0.0, centre - half), min(1.0, centre + half))


@dataclass
class Rate:
    k: int
    n: int

    @property
    def p(self) -> float:
        return self.k / self.n if self.n else math.nan

    @property
    def ci(self) -> tuple[float, float]:
        return wilson(self.k, self.n)

    @property
    def text(self) -> str:
        if not self.n:
            return "no data yet"
        lo, hi = self.ci
        return f"{self.p:.0%} ({self.k}/{self.n}; 95% CI {lo:.0%}–{hi:.0%})"

    @property
    def short(self) -> str:
        return f"{self.p:.0%}" if self.n else "—"


def layer_of(reason: str) -> str:
    if reason.startswith("model:"):
        return "Model flag"
    if reason.startswith("guard:"):
        return "Code guard"
    if reason.startswith("check:"):
        return "Step-6 check"
    return "Stopped early"        # unknown customer, model error, unknown portal


def mask(df: pd.DataFrame, column: str, fn) -> pd.Series:
    """Boolean mask from a per-value test; safe on empty frames and list-valued columns."""
    return df[column].map(lambda v: bool(fn(v))).astype(bool) if len(df) else pd.Series(False, index=df.index, dtype=bool)


def nonempty(df: pd.DataFrame, column: str) -> pd.Series:
    return mask(df, column, lambda v: len(v or []) > 0)


def load_thresholds() -> dict:
    return yaml.safe_load((HERE / "thresholds.yml").read_text(encoding="utf-8"))


# ------------------------------------------------------------------ the evidence
@dataclass
class Evidence:
    lab: object
    frame: pd.DataFrame
    round: int | None
    rounds: list[int]
    workbook: Path
    runs: Path
    demo: bool
    thresholds: dict = field(default_factory=load_thresholds)

    # --- denominators -------------------------------------------------
    @property
    def run(self) -> pd.DataFrame:
        """Cases the workflow actually ran in the app (a recorded routing decision)."""
        return self.frame[self.frame.decision.isin(["send", "escalate"])]

    @property
    def reviewed(self) -> pd.DataFrame:
        return self.frame[self.frame.verdict.isin(["Pass", "Fail"])]

    # --- headline rates -----------------------------------------------
    def pass_rate(self, df: pd.DataFrame | None = None) -> Rate:
        df = self.reviewed if df is None else df[df.verdict.isin(["Pass", "Fail"])]
        return Rate(int((df.verdict == "Pass").sum()), len(df))

    def routing_accuracy(self, df: pd.DataFrame | None = None) -> Rate:
        df = self.run if df is None else df[df.decision.isin(["send", "escalate"])]
        return Rate(int(df.routing_ok.sum()), len(df))

    def missed(self) -> pd.DataFrame:
        r = self.run
        return r[r.expected_human & (r.decision == "send")]

    def unneeded(self) -> pd.DataFrame:
        r = self.run
        return r[~r.expected_human & (r.decision == "escalate")]

    def unsafe_sends(self) -> pd.DataFrame:
        """Auto-sent replies that say a must-not phrase or carry a blocking check flag."""
        r = self.run
        return r[(r.decision == "send") & (nonempty(r, "banned") | nonempty(r, "high_flags"))]

    def code_saves(self) -> pd.DataFrame:
        """Escalations the label required that the model's own flag missed and code caught."""
        r = self.run
        return r[(r.decision == "escalate") & r.expected_human & mask(r, "model_human", lambda v: v is False)]

    def escalation_rate(self) -> Rate:
        r = self.run
        return Rate(int((r.decision == "escalate").sum()), len(r))

    def model_flag_recall(self) -> Rate:
        r = self.run[self.run.expected_human & self.run.model_human.notna()]
        return Rate(int(r.model_human.astype(bool).sum()), len(r))

    # --- breakdowns -------------------------------------------------------
    def by(self, column: str, order: list | None = None) -> pd.DataFrame:
        rows = []
        values = order or sorted(self.frame[column].dropna().unique())
        for v in values:
            part = self.frame[self.frame[column] == v]
            pr, ra = self.pass_rate(part), self.routing_accuracy(part)
            rows.append({column: v, "cases": len(part), "reviewed": pr.n, "pass": pr.k, "pass_rate": pr.p,
                         "pass_lo": pr.ci[0], "pass_hi": pr.ci[1], "run": ra.n, "routing_ok": ra.k,
                         "routing_accuracy": ra.p})
        return pd.DataFrame(rows)

    def layers(self) -> pd.DataFrame:
        """For escalated cases: how often each layer fired, and how often it was the only one."""
        esc = self.run[self.run.decision == "escalate"]
        rows = []
        for layer in LAYERS:
            fired = esc[mask(esc, "layers", lambda ls, lay=layer: lay in ls)]
            alone = fired[mask(fired, "layers", lambda ls: len(ls) == 1)]
            rows.append({"layer": layer, "fired": len(fired), "only_layer": len(alone),
                         "fired_and_needed": int(fired.expected_human.sum())})
        return pd.DataFrame(rows)

    def check_hits(self) -> pd.DataFrame:
        rows = []
        r = self.run
        names = sorted({f for fl in r.all_flags for f in fl})
        for name in names:
            hit = r[mask(r, "all_flags", lambda fl, nm=name: nm in fl)]
            rows.append({"check": name, "cases": len(hit), "auto_sent": int((hit.decision == "send").sum()),
                         "escalated": int((hit.decision == "escalate").sum())})
        return pd.DataFrame(rows, columns=["check", "cases", "auto_sent", "escalated"]).sort_values(
            "cases", ascending=False)

    def disagreements(self) -> pd.DataFrame:
        f, r, rev = self.frame, self.run, self.reviewed
        checked = f[f.expert_check.astype(str).str.len() > 0]
        model_known = r[r.model_human.notna()]
        auto_clean = ~nonempty(rev, "banned") & ~nonempty(rev, "high_flags") & ~nonempty(rev, "missing") \
            & mask(rev, "routing_ok", lambda v: v is not False)
        expert_vs_checks = rev[((rev.verdict == "Pass") & ~auto_clean) | ((rev.verdict == "Fail") & auto_clean)]
        judged = rev[rev.judge_verdict.isin(["Pass", "Fail"])]
        rows = [
            ("Expert vs dataset label", "labels the expert wants changed", len(checked[checked.expert_check != "Agree"]), len(checked)),
            ("Workflow routing vs label", "routed against the label", int(mask(r, "routing_ok", lambda v: v is False).sum()), len(r)),
            ("Model's own flag vs label", "model's human flag against the label", int((model_known.model_human.astype(bool) != model_known.expected_human).sum()), len(model_known)),
            ("Expert vs automatic checks", "pass/fail contradicts the auto checks", len(expert_vs_checks), len(rev)),
            ("Expert vs LLM judge", "judge and expert verdicts differ", int((judged.judge_verdict != judged.verdict).sum()), len(judged)),
        ]
        return pd.DataFrame([{"comparison": a, "what disagrees": b, "k": k, "n": n,
                              "rate": Rate(k, n).short, "ci": Rate(k, n).text} for a, b, k, n in rows])

    def failure_modes(self) -> pd.DataFrame:
        fails = self.reviewed[self.reviewed.verdict == "Fail"]
        counts = fails.failure_mode.replace("", "Not given").fillna("Not given").value_counts()
        return counts.rename_axis("failure mode").reset_index(name="cases")

    def error_cost(self) -> pd.DataFrame:
        cost = self.thresholds["costs_inr"]
        n = len(self.run)
        rows = [("Missed escalation", len(self.missed()), cost["missed_escalation"]),
                ("Unsafe auto-send", len(self.unsafe_sends()), cost["unsafe_auto_send"]),
                ("Unneeded escalation", len(self.unneeded()), cost["unneeded_escalation"])]
        out = pd.DataFrame(rows, columns=["error", "cases", "cost_each_inr"])
        out["cost_inr"] = out.cases * out.cost_each_inr
        out["per_1000_emails_inr"] = (out.cost_inr / n * 1000).round(0) if n else math.nan
        return out

    def gates(self) -> pd.DataFrame:
        g = self.thresholds["gates"]
        pr, ra = self.pass_rate(), self.routing_accuracy()
        rows = [
            ("Expert pass rate (lower end of 95% CI)", f"≥ {g['pass_rate_ci_low']:.0%}",
             f"{pr.ci[0]:.0%}" if pr.n else "—", pr.n > 0 and pr.ci[0] >= g["pass_rate_ci_low"]),
            ("Routing accuracy", f"≥ {g['routing_accuracy']:.0%}", ra.short,
             ra.n > 0 and ra.p >= g["routing_accuracy"]),
            ("Missed escalations", f"≤ {g['missed_escalations']}", str(len(self.missed())) if ra.n else "—",
             ra.n > 0 and len(self.missed()) <= g["missed_escalations"]),
            ("Unsafe auto-sends", f"≤ {g['unsafe_auto_sends']}", str(len(self.unsafe_sends())) if ra.n else "—",
             ra.n > 0 and len(self.unsafe_sends()) <= g["unsafe_auto_sends"]),
            ("Cases run in this round", f"= {len(self.frame)}", str(ra.n), ra.n == len(self.frame)),
        ]
        return pd.DataFrame(rows, columns=["gate", "target", "actual", "met"])

    def provenance(self) -> dict:
        r = self.run
        models = sorted({f"{p} {m}".strip() for p, m in zip(r.provider.fillna(""), r.model.fillna("")) if p})
        return {"round": self.round, "rounds in file": ", ".join(map(str, self.rounds)) or "none",
                "drafting model": ", ".join(models) or "—", "workbook": self.workbook.name,
                "workbook sha256": sha256(self.workbook)[:16], "runs file": str(self.runs.name),
                "runs sha256": sha256(self.runs)[:16], "demo data": "yes" if self.demo else "no"}


def build_frame(lab, records: dict[str, dict]) -> pd.DataFrame:
    rows = []
    for case in lab.cases:
        rec = records.get(case["id"], {})
        ev = eval_checks(case, rec)
        flags = rec.get("flags") or []
        reasons = rec.get("reasons") or []
        decision = rec.get("decision")
        rows.append({
            "id": case["id"], "customer_id": case["customer_id"], "message": case["message"],
            "intent": case["expected_intent"], "rule": case["rule"], "expected_human": bool(case["needs_human"]),
            "tags": case["tags"], "language": case["language"], "tone": case["tone"],
            "difficulty": case["difficulty"], "kaggle": case["kaggle"], "expert_check": case.get("expert_check") or "",
            "expert_note": case.get("expert_note") or "", "rationale": case["rationale"],
            "decision": decision, "model_human": rec.get("model_needs_human"), "model_intent": rec.get("model_intent"),
            "outcome": rec.get("outcome"), "reply": rec.get("reply") or "", "reasons": reasons,
            "layers": sorted({layer_of(r) for r in reasons}) if decision == "escalate" else [],
            "verdict": rec.get("verdict") or None, "failure_mode": rec.get("failure_mode") or "",
            "critique": rec.get("critique") or "", "judge_verdict": rec.get("judge_verdict") or None,
            "routing_ok": ev["routing_ok"] if decision else None, "intent_ok": ev["intent_ok"],
            "missing": ev["missing"], "banned": ev["banned"],
            "high_flags": [f["check"] for f in flags if f.get("severity") == "high"],
            "all_flags": [f["check"] for f in flags],
            "provider": rec.get("provider"), "model": rec.get("model"), "latency_s": rec.get("latency_s"),
            "cost_usd": rec.get("cost_usd"), "error": rec.get("error"),
        })
    return pd.DataFrame(rows)


def load(review_round: int | None = None) -> Evidence:
    """Load the workbook and the saved runs for one round (default: the latest round in the file)."""
    lab = load_lab(WORKBOOK)
    raw = json.loads(RUNS.read_text(encoding="utf-8")) if RUNS.exists() else {}
    rounds = sorted({int(v["round"]) for v in raw.values()})
    rnd = int(review_round) if review_round else (rounds[-1] if rounds else None)
    records = {v["case_id"]: v for v in raw.values() if rnd is not None and int(v["round"]) == rnd}
    demo = any(v.get("verdict_source") == "demo-auto" for v in records.values())
    return Evidence(lab=lab, frame=build_frame(lab, records), round=rnd, rounds=rounds, workbook=WORKBOOK,
                    runs=RUNS, demo=demo)


# ------------------------------------------------------------------ workbook-only views (dataset card)
def workbook_sheet(name: str) -> pd.DataFrame:
    ws = load_workbook(WORKBOOK, data_only=True)[name]
    rows = list(ws.iter_rows(values_only=True))
    header = [str(h) if h is not None else f"col{i}" for i, h in enumerate(rows[0])]
    return pd.DataFrame([r for r in rows[1:] if r and r[0]], columns=header)


# ------------------------------------------------------------------ presentation helpers
def plotly_layout(fig, height: int = 320):
    fig.update_layout(height=height, margin=dict(l=8, r=28, t=8, b=8), plot_bgcolor="#fcfcfb",
                      paper_bgcolor="rgba(0,0,0,0)", font=dict(family="Inter, system-ui, sans-serif", size=13, color=INK),
                      hoverlabel=dict(bgcolor="white", font_size=13), showlegend=False)
    fig.update_xaxes(gridcolor="#E7ECF2", zeroline=False, linecolor="#D6E0EC", tickfont=dict(color=MUTED))
    fig.update_yaxes(gridcolor="#E7ECF2", zeroline=False, linecolor="#D6E0EC", tickfont=dict(color=MUTED))
    return fig


def status(met: bool) -> str:
    return (f'<span style="color:{GOOD};font-weight:600">✓ met</span>' if met
            else f'<span style="color:{CRITICAL};font-weight:600">✗ not met</span>')


def tiles(items: list[tuple[str, str, str]]) -> str:
    """KPI tiles: (label, value, note)."""
    cells = "".join(f'<div class="kpi"><div class="kpi-label">{a}</div><div class="kpi-value">{b}</div>'
                    f'<div class="kpi-note">{c}</div></div>' for a, b, c in items)
    return f'<div class="kpis">{cells}</div>'


def esc(text) -> str:
    import html
    return html.escape(str(text or ""))


def html_table(df: pd.DataFrame, columns: dict[str, str] | None = None, escape: bool = True) -> str:
    """A compact Bootstrap table; `columns` maps source column -> header text (and picks the order)."""
    if columns:
        df = df[list(columns)].rename(columns=columns)
    return df.to_html(index=False, escape=escape, border=0, na_rep="—",
                      classes="table table-sm table-hover compact")


def banner(ev: "Evidence") -> str:
    """Markdown callouts explaining what the data can and can't say. Print with `output: asis`."""
    out = []
    if ev.demo:
        out.append("::: {.callout-warning}\n## Demo data\nDrafts come from the offline stand-in and verdicts were "
                   "derived from the automatic checks, not from an expert. Use this to preview the report, "
                   "never as evidence.\n:::")
    if ev.round is None:
        out.append("::: {.callout-note}\n## No runs yet\nRun cases in Acme Review Studio, review them, then render "
                   "again. The numbers below fill in as runs and verdicts are saved.\n:::")
    elif len(ev.run) < len(ev.frame) or len(ev.reviewed) < len(ev.frame):
        out.append(f"::: {{.callout-note}}\n## Partial round\n{len(ev.run)} of {len(ev.frame)} cases have been run "
                   f"and {len(ev.reviewed)} reviewed. Every rate shows its denominator; read small ones with care.\n:::")
    return "\n\n".join(out)


_PLOTLY_INCLUDED = False


def fig_html(fig, height: int | None = None) -> str:
    """Plotly figure as HTML. The first chart on a page loads plotly.min.js, which the project's
    pre-render step copies next to the pages, so charts work offline and need no CDN."""
    global _PLOTLY_INCLUDED
    include = "directory" if not _PLOTLY_INCLUDED else False
    _PLOTLY_INCLUDED = True
    return fig.to_html(full_html=False, include_plotlyjs=include, default_height=height or fig.layout.height or 320,
                       config={"displayModeBar": False, "responsive": True})
