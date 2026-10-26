"""Copy the Acme lab's new cases, customers and policy decisions into the studio workbook.

Reads fde-first-principles/data/acme (cases_abuse.jsonl, cases_multiturn.jsonl, customers.json,
policy.md) and adds only what the workbook doesn't have yet. Existing rows, expert columns and
reviews are left alone, so it is safe to re-run after the jsonl files change.

    python import_cases.py
    ACME_XLSX=path/to/other.xlsx python import_cases.py
"""
from __future__ import annotations

import json
import os
import re
from copy import copy
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from acme.lab import CUSTOMER_HISTORY, _restore_chart

ROOT = Path(__file__).parent
XLSX = Path(os.environ.get("ACME_XLSX", ROOT / "data" / "Acme_Fibre_Grounded_Draft_Eval.xlsx"))
LAB_DATA = ROOT.parents[2] / "data" / "acme"
CASE_FILES = ["cases_abuse.jsonl", "cases_multiturn.jsonl"]
CASE_COLUMNS = ["history", "concession_allowed", "pair_id", "twin"]
LANGUAGES = "English,Hinglish,Hindi (Devanagari),Marathi-English"
MANIPULATIVE = {"injection", "fake-policy", "fake-offer", "fake-code", "fake-memory", "forged-approval",
                "insider-claim"}


def _tone(tags: list[str]) -> str:
    tags = set(tags)
    if tags & MANIPULATIVE:
        return "Manipulative"
    if "angry" in tags:
        return "Angry"
    if tags & {"urgent", "pressure"}:
        return "Urgent"
    return "Neutral"


def _copy_style(ws, src_row: int, dst_row: int, columns: range):
    for col in columns:
        ws.cell(row=dst_row, column=col)._style = copy(ws.cell(row=src_row, column=col)._style)


def _add_headers(ws, names: list[str], like_column: int) -> dict[str, int]:
    col = {c.value: c.column for c in ws[1] if c.value}
    for name in names:
        if name not in col:
            cell = ws.cell(row=1, column=ws.max_column + 1, value=name)
            cell._style = copy(ws.cell(row=1, column=like_column)._style)
            col[name] = cell.column
    return col


def _last_row(ws, pattern: str) -> int:
    return max(r for r in range(1, ws.max_row + 1) if re.match(pattern, str(ws.cell(row=r, column=1).value or "")))


def _take_note(ws, prefix: str) -> str | None:
    """Remove the 'Source: ...' footnote so appended rows don't overwrite it; put it back with _put_note."""
    note = None
    for r in range(2, ws.max_row + 1):
        value = ws.cell(row=r, column=1).value
        if isinstance(value, str) and value.startswith(prefix):
            note = value
            ws.cell(row=r, column=1).value = None
    return note


def _put_note(ws, last: int, note: str | None) -> None:
    if note:
        ws.cell(row=last + 2, column=1).value = note


def import_cases(wb) -> int:
    ds = wb["Dataset"]
    col = _add_headers(ds, CASE_COLUMNS, like_column=ds.max_column)
    existing = {str(ds.cell(row=r, column=1).value) for r in range(2, ds.max_row + 1)}
    last = _last_row(ds, r"^[a-z]\d+")
    crm = ds.cell(row=2, column=col["CRM snapshot"]).value
    added = 0
    for name in CASE_FILES:
        for line in (LAB_DATA / name).read_text(encoding="utf-8").splitlines():
            if not line.strip():
                continue
            case = json.loads(line)
            if case["id"] in existing:
                continue
            last += 1
            values = {
                "id": case["id"], "customer_id": case["customer_id"], "CRM snapshot": re.sub(r"\$B2\b", f"$B{last}", crm),
                "message": case["message"], "expected_intent": case["expected_intent"],
                "needs_human": bool(case["needs_human"]), "escalation_rule": case["escalation_rule"],
                "must_include": " | ".join(case["must_include"]), "must_not": " | ".join(case["must_not"]),
                "kb_sources": ", ".join(case["kb_sources"]), "tags": ", ".join(case["tags"]),
                "language": case["language"], "tone": _tone(case["tags"]), "difficulty": case["difficulty"],
                "rationale": case["rationale"],
                "history": json.dumps(case["history"], ensure_ascii=False) if case.get("history") else None,
                "concession_allowed": case.get("concession_allowed"),
                "pair_id": case.get("pair_id"), "twin": case.get("twin"),
            }
            _copy_style(ds, 2, last, range(1, ds.max_column + 1))
            for key, value in values.items():
                ds.cell(row=last, column=col[key]).value = value
            existing.add(case["id"])
            added += 1
    ds.auto_filter.ref = f"A1:{get_column_letter(ds.max_column)}{last}"
    for dv in ds.data_validations.dataValidation:
        if "L2" in str(dv.sqref):
            dv.formula1 = f'"{LANGUAGES}"'
    return added


def import_customers(wb) -> int:
    cs = wb["Customers"]
    note = _take_note(cs, "Source:")
    col = _add_headers(cs, CUSTOMER_HISTORY, like_column=8)
    rows = {cs.cell(row=r, column=1).value: r for r in range(2, cs.max_row + 1)
            if re.match(r"^AC\d+$", str(cs.cell(row=r, column=1).value or ""))}
    last = max(rows.values())
    added = 0
    for c in json.loads((LAB_DATA / "customers.json").read_text(encoding="utf-8")):
        r = rows.get(c["id"])
        if r is None:
            last += 1
            r = last
            _copy_style(cs, 2, r, range(1, cs.max_column + 1))
            for k, key in enumerate(["id", "name", "area", "plan", "status", "tenure_months", "balance_due",
                                     "complaints_last_30_days"], start=1):
                cs.cell(row=r, column=k).value = c[key]
            added += 1
        history = dict(c, payments_this_cycle=", ".join(map(str, c["payments_this_cycle"])) or None)
        for name in CUSTOMER_HISTORY:
            if cs.cell(row=r, column=col[name]).value in (None, ""):
                cs.cell(row=r, column=col[name]).value = history.get(name)
    _put_note(cs, last, note)
    return added


def import_policy(wb) -> int:
    pk = wb["Policy_KB"]
    note = _take_note(pk, "Source:")
    have = {pk.cell(row=r, column=1).value for r in range(2, pk.max_row + 1)}
    text = (LAB_DATA / "policy.md").read_text(encoding="utf-8")
    rules = re.findall(r"^(\d+)\. (.+?)(?=^\d+\. |^## |\Z)", text, flags=re.M | re.S)
    last = _last_row(pk, r"^(KB|P)\d+$")
    added = 0
    for number, body in rules:
        rid = f"P{number}"
        if rid in have:
            continue
        last += 1
        _copy_style(pk, last - 1, last, range(1, 4))
        for k, value in enumerate([rid, "Reply policy", " ".join(body.split())], start=1):
            pk.cell(row=last, column=k).value = value
        added += 1
    _put_note(pk, last, note)
    return added


def main():
    wb = load_workbook(XLSX)
    counts = {"cases": import_cases(wb), "customers": import_customers(wb), "policy rules": import_policy(wb)}
    _restore_chart(wb)
    wb.save(XLSX)
    print(f"{XLSX.name}: added " + ", ".join(f"{n} {what}" for what, n in counts.items()))


if __name__ == "__main__":
    main()
